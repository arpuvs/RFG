import sys, visa, time
# sys.path.append('../../Common/FMB_USB_Python_Files')
sys.path.append('../../Common')
# # C:\Git_Repositories\RFG\Common\FMB_USB_Python_Files
# from ADI_GPIB.AgilentN5181A import *
from ADI_GPIB.AgilentN9030A import *
# # from ADI_GPIB.AgilentN6705B import *
# from FMB import *
#
# import xlsxwriter
import openpyxl
from openpyxl import *
from openpyxl.utils import *
# import atexit
#
# def exit_handler():
#     print 'Handler run'
#
# atexit.register(exit_handler)

# asdfsadf asdfasd

# fmbDict = {1: "2.5 MHz", 2: "5 MHz", 3: "33 MHz", 4: "78 MHz",
#           5: "120 MHz", 6: "225.3 MHz", 7: "350.3 MHz", 8: "500 MHz",
#           9: "800.3 MHz", 10: "1 GHz", 11: "1.5 GHz", 12: "2 GHz",
#           13: "2.5 GHz", 14: "3 GHz", 15: "3.5 GHz", 16: "4 GHz",
#           17: "4.5 GHz", 18: "5 GHz", 19: "5.5 GHz", 20: "5.9 GHz",
#           21: "Aux"}
#

# fh = open('Test.csv', 'w')

wb = load_workbook(filename='C:\\Users\\bsulliv2\\Desktop\\Pronghorn_Results\\Test\\OutputTest.xlsx')
sheet_ranges = wb['Sheet1']
# print sheet_ranges
print sheet_ranges['A']

A = sheet_ranges['A']
# print AA
val = A[0].value
print val
#
# print len(AA)
freq = [1e9, 2e9, 3e9, 4e9, 5e9, 6e9, 7e9, 8e9, 9e9]
data1 = [1, 2, 3, 4, 5, 6, 7, 8, 9]
data2 = [10, 20, 30, 40, 50, 60, 70, 80, 90]

datalabel = ['DataA', 'Data1_A']

for row in range(len(A)+1, len(A) + len(data1) + 1):
    # for col in range(1, 10):
    for i in range(len(datalabel)):
        sheet_ranges.cell(column=i+1, row=row, value=datalabel[i])
    sheet_ranges.cell(column=i+2, row=row, value=freq[row - (len(A) + 1)])
    sheet_ranges.cell(column=i+3, row=row, value=data1[row - (len(A) + 1)])

datalabel[1] = 'Data1_B'
A = sheet_ranges['A']

for row in range(len(A)+1, len(A) + len(data2) + 1):
    # for col in range(1, 10):
    for i in range(len(datalabel)):
        sheet_ranges.cell(column=i+1, row=row, value=datalabel[i])
    sheet_ranges.cell(column=i+2, row=row, value=freq[row - (len(A) + 1)])
    sheet_ranges.cell(column=i+3, row=row, value=data2[row - (len(A) + 1)])
        # sheet_ranges.cell(column=2, row=row, value='Data1A')
        # sheet_ranges.cell(column=3, row=row, value='Data1')

wb.save(filename='C:\\Users\\bsulliv2\\Desktop\\Pronghorn_Results\\Test\\OutputTest.xlsx')
# wb = Workbook()
#
# dest_filename = 'openpyxlTest.xlsx'
#
# ws1 = wb.active
# ws1.title = 'Worksheet title'
#
# for row in range(1, 40):
#     ws1.append(range(600))
#
# ws2 = wb.create_sheet(title='WS2')
# ws2['F5'] = 3.14
#
# ws3 = wb.create_sheet(title='Worksheet 3')
# for row in range(10,20):
#     for col in range(27,54):
#         _ = ws3.cell(column = col, row=row, value='{0}'.format(get_column_letter(col)))
# print(ws3['AA10'].value)
#
# wb.save(filename=dest_filename)


# workbook = xlsxwriter.Workbook('Test.xlsx')
# worksheet = workbook.add_worksheet()
# meas1 = range(1, 101)
# meas2 = range(10, 1001, 10)
# meas3 = range(100, 10001, 100)
#
# col = 0
# row = 0
# for i in range(len(meas1)):
#     row = i
#     worksheet.write(row, col, meas1[i])
#
# col = col + 1
# for i in range(len(meas2)):
#     row = i
#     worksheet.write(row, col, meas2[i])
#
# col = col + 1
# for i in range(len(meas3)):
#     row = i
#     worksheet.write(row, col, meas3[i])

# workbook.close()
# for i in range(len(meas1)):
#     line = [meas1[i], meas2[i], meas3[i]]
#     fh.write(str(line).strip('[]'))
#     fh.write('\n')
# fh.write(str(meas1).strip('[]'))
# fh.write('\n')
# fh.write(str(meas2).strip('[]'))
# fh.write('\n')
# fh.write(str(meas3).strip('[]'))
# fh.write('\n')


# print meas1, meas2, meas3
# Filter = FMB('COM3', fmbDict)
# Source2 = AgilentN5181A(11)
# Analyzer = AgilentN9030A(18)
#
# Analyzer.__SetAutoAtten__(1)
# # Analyzer.instr.write(':POW:ATTN:AUTO 1')
# time.sleep(1)
# Analyzer.__SetAutoAtten__(0)
# Analyzer.__SetAtten__(12)
#
# print Analyzer.__GetAtten__()
# # # # Filter.select_filter(3)
# # print Filter.get_filter_center_freq_hz(19)
#
# # Filter box test
# for i in range(1, 21):
#     # Filter.select_filter(i)
#     # Source2.__SetFreq__
#     val = fmbDict[i].split()
#     if val[1] == 'MHz': val = float(val[0]) * 1e6
#     elif val[1] == 'GHz': val = float(val[0]) * 1e9
#     print val
#     Source2.__SetFreq__(val)
#     Analyzer.__Setfc__(val)
#     Filter.select_filter(i)
#     Analyzer.__SetMarkerFreq__(1, val)
#     time.sleep(0.5)
#     carrierMag = float(Analyzer.__GetMarkerAmp__(1))

# Want -2dBm out of part

# Supply = AgilentN6705B(26)
# Source1 = AgilentN5181A(20)
# Source2 = AgilentN5181A(11)
# Analyzer = AgilentN9030A(18)
# Oven = WatlowF4(4)
# # #
# #
# print Oven.__GetTemp__()
# print Oven.__GetSP__()
# from ADI_GPIB.WatlowF4 import *
# from ADI_GPIB.AgilentE5071C import *
# from ADI_GPIB.KeysightN5242A import *
#
# # VNA = AgilentE5071C(17)
# # Oven = WatlowF4(4)
# VNA = KeysightN5424A(17)
#
# VNA.__LoadState__('default.csa')
# VNA.__SetSweepType__(1, 'LOG')
# VNA.__SetStartf__(1, 10e6)
# VNA.__SetStopf__(1, 10.01e9)
# VNA.__SetNumPoints__(1, 1e3)
# VNA.__SetAvg__(1, 1)
# VNA.__SetAutoTime__(1, True)
# VNA.__SetTrigType__('MAN')
# VNA.__SetContinuous__(1, False)
# VNA.__EnableAvg__(1, True)
# # VNA.__EnableTrigAvg__(True)
# VNA.__SetTopology__(1, 'BBAL')
# VNA.__SetPorts__(1, 1, 3, 2, 4)
# VNA.__EnableBal__(1)
# VNA.instr.write('SENS:CORR:CSET:ACT \"BS_Cal\", 1')
# VNA.__SetActiveTrace__(1, 1)
#
# #
# VNA.__AddWindow__(2)
# VNA.__AddMeas__(2, 'P1dB', 'Gain Compression')
# VNA.__AddTrace__(2, 1, 'P1dB')
# VNA.__SetSweepType__(2, 'LOG')
# VNA.__SetStartf__(2, 10e6)
# VNA.__SetStopf__(2, 10.01e9)
# VNA.__SetNumPoints__(2, 1e3)
# VNA.__SetAvg__(2, 1)
# VNA.__SetAutoTime__(2, True)
# # VNA.__SetTrigType__('MAN')
# VNA.__SetContinuous__(2, False)
# VNA.__EnableAvg__(2, True)
# # VNA.__EnableTrigAvg__(True)
# VNA.__SetTopology__(2, 'BBAL')
# VNA.__SetPorts__(2, 1, 3, 2, 4)
# VNA.__EnableBal__(2)
#
# VNA.__AddWindow__(3)
# VNA.__AddMeas__(3, 'NF', 'Noise Figure Cold Source')
# VNA.__AddTrace__(3, 1, 'NF')
# VNA.__SetStartf__(3, 10e6)
# VNA.__SetStopf__(3, 10.01e9)
# VNA.__SetNumPoints__(3, 1e3)
# VNA.__SetAvg__(3, 1)
# VNA.__SetAutoTime__(3, True)
# VNA.__SetTrigType__('MAN')
# VNA.__SetContinuous__(3, False)
# VNA.__EnableAvg__(3, True)
# # VNA.__EnableTrigAvg__(True)
# VNA.__SetTopology__(3, 'BBAL')
# VNA.__SetPorts__(3, 1, 3, 2, 4)
# VNA.__EnableBal__(3)
#
# VNA.__AddWindow__(4)
# VNA.__AddMeas__(4, 'IMD', 'Swept IMD')
# VNA.__AddTrace__(4, 1, 'IMD')
# VNA.__IMDDelta__(4, 2e6)    # Undefined header
# VNA.__IMDStart__(4, 10e6)   # Parameter not valid
# VNA.__IMDStop__(4, 10.01e9) # Parameter not valid
# VNA.__SetStartf__(4, 10e6)
# VNA.__SetStopf__(4, 10.01e9)
# VNA.__SetNumPoints__(4, 1e3)
# VNA.__SetAvg__(4, 1)
# VNA.__SetAutoTime__(4, True)
# VNA.__SetTrigType__('MAN')
# VNA.__SetContinuous__(4, False)
# VNA.__EnableAvg__(4, True)
# # VNA.__EnableTrigAvg__(True)
# VNA.__SetTopology__(4, 'BBAL')
# VNA.__SetPorts__(4, 1, 3, 2, 4)
# VNA.__EnableBal__(4)

# VNA.instr.write('CALC1:CUST:DEF \'P1dB\', \'Gain Compression\' ')
# def meas():
#     VNA.__InitMeas__(1)
#     VNA.__SingleTrig__()    # Causes error - Unidentified header
#     status = 0
#     while status == 0:
#         status = VNA.__CheckStatus__()
#         time.sleep(0.1)
#     ans = VNA.__GetData__(1)    # Causes error
#     print ans
#     ans = ans.split(',')
#     magans = []
#     imagans = []
#     for val in range(len(ans)):
#         if (val % 2) == 0:
#             magans.append(float(ans[val]))
#         else:
#             imagans.append(float(ans[val]))
#
#     compans = []
#     for val in range(len(magans)):
#         compans.append(magans[val] + imagans[val] * 1j)
#
#     # return compans
#     return magans, imagans
#
# VNA.__SetSweepType__(1, 'LOG')
# VNA.__SetStartf__(1, 10e06)
# VNA.__SetStopf__(1, 10.01e09)
# VNA.__SetNumPoints__(1, 1e3)
# VNA.__SetAvg__(1, 16)
# VNA.__SetAutoTime__(1, True)
# VNA.__SetTrigType__('MAN')
# VNA.__SetContinuous__(1, False)
# VNA.__EnableAvg__(1, True)
# # VNA.__EnableTrigAvg__(True)     # Causes error - Unidentified header
#
# VNA.__SetActiveTrace__(1, 1)
# VNA.__SetBalBal__(1, 'Bal-Bal')
# VNA.__SetBBalParam__(1, 1, 'SCC21')
# VNA.__SetActiveFormat__(1, 'MLOG')
# scc21_mlog = meas()

# VNA.__SetTopology__(1, 'BBAL')
# VNA.__SetBBalParam__(1, 1, 'SCC21')
# VNA.__SetBBalParam__(1, 1, 'SDD11')

# VNA.instr.write('SENS:CORR:CSET:ACT \"BS_Cal\", 1')

# VNA.instr.write('CALC1:FSIM:BAL:PAR:BBAL SDD11')


# print VNA.instr.ask('CALC:PAR:CAT?')
# VNA.__SetActiveTrace__(1, 1)    # Causes error - Missing parameter
# ans = VNA.__GetData__(1)    # Causes error
# print ans


# print Analyzer.__ClearAverage__()
# print Analyzer.__SetAverage__(50)
# print time.sleep(1)
# print Analyzer.__CheckStatus__(600)

# print Analyzer.__Align__()
# print Analyzer.__CheckStatus__(600)
# print Analyzer.__CheckStatus__(30)
# Analyzer.__Align__()
# print Analyzer.__CheckStatus__(300)


# date = time.ctime(time.time())
# date = date.replace(':', '.')
# fh = open('Intermod_Dist_' + date + '.csv', 'w')
#
# # Source1.__SetState__(1)
# # Source2.__SetState__(1)
# # print Source1.__GetState__()
# # print Source2.__GetState__()
# # time.sleep(0.5)
# # Source1.__SetState__(0)
# # Source2.__SetState__(0)
# # print Source1.__GetState__()
# # print Source2.__GetState__()
#
# # Analyzer.__StoreState__('BS_Test_State')
# # time.sleep(1)
# # Analyzer.__LoadState__('BS_Test_State')
# # Source1.__SetFreq__(3.999e9)
# # Source2.__SetFreq__(4.001)
# # print Source1.__GetFreq__()
#
# # print Analyzer.__Setfc__(4.03e9)
# # Analyzer.__SetMarkerFreq__(1, 4.003e9)
# # print Analyzer.__GetMarkerAmp__(1)
# # Analyzer.__MarkerMax__(1)
# # print Analyzer.__GetSpan__()
# # print Analyzer.__SetSpan__(10e9)
# # print Analyzer.__Setfc__(4.03e9)
# # Analyzer.__LoadState__('BS_Test_State')
#
# # Analyzer.__LoadState__('BS_Test_State')
#
# freqlist = [100e6, 250e6, 500e6, 1e9, 4e9]
# lowerPeak = []
# higherPeak = []
#
# Analyzer.__SetSpan__(10e3)
# for i in range(11):
#     for freq in freqlist:
#         Source1.__SetFreq__(freq-1e6)
#         Source2.__SetFreq__(freq+1e6)
#         Analyzer.__Setfc__(freq-3e6)
#         Analyzer.__SetMarkerFreq__(1, freq-3e6)
#         time.sleep(1)
#         lowerPeak.append(float(Analyzer.__GetMarkerAmp__(1)))
#         Analyzer.__Setfc__(freq+3e6)
#         Analyzer.__SetMarkerFreq__(1, freq+3e6)
#         time.sleep(1)
#         higherPeak.append(float(Analyzer.__GetMarkerAmp__(1)))
#
#
#     fh.write('Test %d' % i)
#     fh.write('\n')
#     fh.write(str(freqlist).strip('[]'))
#     fh.write('\n')
#     fh.write(str(lowerPeak).strip('[]'))
#     fh.write('\n')
#     fh.write(str(higherPeak).strip('[]'))
# Analyzer.__SetAverage__(100e3)
# print Analyzer.__GetAverage__()
# print Analyzer.__GetAverage__()
# print Analyzer.__GetAverage__()
# print Analyzer.__GetAverage__()
# Analyzer.__ClearAverage__()
# for Temp in range(65496, 65496 + 100):
# print Oven.__SetTemp__(25)
# print Oven.__GetTemp__()
# Supply.__Enable__(0, 1)
# print Supply.__SetI__(0.04, (1, 2))
# print Supply.__SetV__(0, 1)
#
# print 1
# print 2
# raise Exception ('Blah')
# print 3

# print Supply.__GetI__(1)
